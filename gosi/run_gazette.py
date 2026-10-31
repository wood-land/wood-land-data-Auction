"""
run_gazette.py — 지자체 공보(시보·군보·도보·구보·공보) 최근 N일 게시물에서 보상 관련 주소 추출.

사용 예
  python run_gazette.py                          # 엑셀의 모든 공보 게시판, 최근 15일
  python run_gazette.py --city 청주시 --city 평창군   # 일부 지자체만
  python run_gazette.py --days 30 --workers 6

게시판 목록의 원본은 엑셀(output/시군구_공보게시판_목록.xlsx, '공보게시판' 시트)이다.
  - 'D열 종류 / E열 게시판 URL / F열 판정' 을 읽고, 'L열 사용자 입력 URL'이 있으면 그것을 우선한다.
  - 판정이 '미발견'이고 사용자 입력 URL도 없으면 건너뛴다(수동 확인 필요 목록).
  - 판정이 '제외'로 시작하면 URL 이 있어도 수집하지 않는다(사용자가 뺀 게시판: 경기도청 도보는 양이 너무 많고, 서산시·제주시·서귀포시는 공보가 없다).
처리: 목록(최근 N일) -> 게시글 첨부 -> 미리보기(Synap/한컴 뷰어) OCR 우선, 안 되면 다운로드(메모리) 후
본문 추출/스캔 PDF OCR -> 공고 단위 분할 -> 보상 관련 공고만 -> 소재지·지번 추출(core/gazette_extractor.py).
결과: output/gazette/gazette_addresses.csv(주소 행), gazette_notices.csv(보상 관련 공고 목록),
      gazette_status.json(게시판별 처리 결과). 엑셀 갱신은 별도 스크립트가 status를 읽어 반영한다.
"""
import argparse
import csv
import datetime as dt
import hashlib
import importlib
import io
import zipfile
import json
import os
import sys
import time
import traceback
import warnings
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import requests  # noqa: E402
import ssl  # noqa: E402
from requests.adapters import HTTPAdapter  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core import gazette_extractor as GX  # noqa: E402
from core import http_fetch  # noqa: E402
from core import synap_preview as SP  # noqa: E402
from core.sqlite_store import post_key  # noqa: E402
from run_gosigongo import discover_provinces  # noqa: E402

warnings.filterwarnings("ignore")
DEFAULT_EXCEL = os.path.join(HERE, "output", "시군구_공보게시판_목록.xlsx")
if not os.path.exists(DEFAULT_EXCEL):  # git에는 _new 파일이 최신본으로 올라가 있다
    DEFAULT_EXCEL = os.path.join(HERE, "output", "시군구_공보게시판_목록_new.xlsx")
MAX_PAGES = 300  # 안전장치일 뿐: 보통은 N일보다 오래된 글을 만나거나 목록이 끝나면 먼저 멈춘다
MAX_ATTACH_PER_POST = 30
UNDATED_TAKE = 3
MAX_BYTES = 300 * 1024 * 1024
DOC_EXT = {"pdf", "hwpx", "hwp", "docx", "doc", "zip", ""}


class LegacyTLSAdapter(HTTPAdapter):
    """구형 TLS/약한 암호 서버(SSLError: handshake 실패)도 읽는다. 공개 게시판 읽기 전용이라 인증서 검증은 끈다."""

    def init_poolmanager(self, *args, **kwargs):
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        try:
            ctx.set_ciphers("DEFAULT@SECLEVEL=0")
        except ssl.SSLError:
            pass
        ctx.options |= getattr(ssl, "OP_LEGACY_SERVER_CONNECT", 0x4)
        kwargs["ssl_context"] = ctx
        super().init_poolmanager(*args, **kwargs)


def new_session():
    s = requests.Session()
    s.verify = False
    s.mount("https://", LegacyTLSAdapter())
    return s


def load_boards(excel_path):
    wb = load_workbook(excel_path, data_only=True)
    ws = wb["공보게시판"]
    heads = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(heads)}
    boards = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["지자체"]]:
            continue
        user_url = row[idx["사용자 입력 URL"]] if "사용자 입력 URL" in idx else None
        url = (user_url or "").strip() if user_url else ""
        judged = row[idx["판정"]]
        if str(judged or "").startswith("제외"):
            continue
        if not url and judged != "미발견":
            url = row[idx["게시판 URL"]] or ""
        if not url:
            continue
        boards.append({
            "module": row[idx["권역(파일)"]], "province": row[idx["광역 단위"]], "city": row[idx["지자체"]],
            "kind": row[idx["종류"]] or "공보", "url": url, "judged": judged, "manual": bool(user_url),
        })
    return boards


def regions_index():
    P = discover_provinces()
    return {(m, c.city): c for m, mod in P.items() for c in mod.CITIES}


def download(session, url, referer=None):
    headers = dict(UA_HEADERS)  # 짧은 UA("Mozilla/5.0")는 일부 사이트(대전 서구 등)가 404로 막는다
    if referer:  # 일부 사이트(당진시 등)는 Referer가 없으면 파일 대신 오류 HTML을 준다
        headers["Referer"] = referer
    def once():
        r = session.get(url, headers=headers, timeout=90, stream=True)
        r.raise_for_status()
        ctype = r.headers.get("content-type", "")
        buf = bytearray()
        for chunk in r.iter_content(1 << 16):
            buf += chunk
            if len(buf) > MAX_BYTES:
                raise ValueError("파일이 너무 큼")
        return bytes(buf), ctype

    return http_fetch.retry_call(once)  # 일시 오류(400/5xx/시간초과)는 다시 받는다(2026-10-25: 밀양시 공보 첫 시도 실패, 재시도 성공)


DUP = object()  # 같은 파일이 이름 링크/다운로드 링크로 두 번 잡힌 경우의 표식
UA_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"}


def _short(text):
    return not text or len(text.strip()) < 200


class OcrText(str):
    """OCR 로 읽은 본문(글자 층이 없는 이미지 문서·미리보기 이미지). 여기서 나온 필지는 숫자 오인식 위험이 있어 parcels.ocr=1 로 표시한다."""


def read_attachment_text(session, att, referer, seen_hash, stat):
    """첨부 1개 -> 본문 텍스트. 순서: 미리보기(Synap/한컴 뷰어) 페이지 이미지 OCR을 먼저 시도하고,
    미리보기가 없거나 실패하면 다운로드해 직접 읽기 -> 스캔형 PDF는 OCR. 중복 파일이면 DUP, 읽을 수 없으면 None."""
    if att.get("preview_url"):
        try:
            t = SP.preview_text(session, att["preview_url"], referer, UA_HEADERS)
        except (requests.RequestException, ValueError, OSError):
            t = None
        if not _short(t):
            h = hashlib.sha1(t.encode("utf-8", "ignore")).hexdigest()
            if h in seen_hash:
                return DUP
            seen_hash.add(h)
            stat["preview"] += 1
            return OcrText(t)
    text, kind, data = None, None, None
    try:
        data, _ = download(session, att["url"], referer=referer)
        h = hashlib.sha1(data).hexdigest()
        if h in seen_hash:
            return DUP
        seen_hash.add(h)
        kind = GX.sniff_kind(data)
    except (requests.RequestException, ValueError):
        data = None
    if data is not None and kind in ("pdf", "hwpx", "docx", "hwp"):
        text = GX.document_text(data, kind)
        if kind == "hwp" and not text:  # 암호/배포용 등 읽을 수 없는 한글 문서
            stat["hwp_only"] += 1
        if kind == "pdf" and _short(text) and len(data) > 200_000 and GX.ocr_available():
            text = GX.ocr_pdf_text(data)  # 스캔형(이미지) PDF -> OCR
            stat["ocr"] += 1
            if text:
                text = OcrText(text)
    elif data is not None and kind == "zip":  # 여러 파일을 묶어 내려주는 사이트(상주시 등): 안의 문서를 모두 읽는다
        parts, via_ocr = [], False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for name in z.namelist():
                    if name.lower().endswith((".pdf", ".hwpx", ".docx", ".hwp")):
                        member = z.read(name)
                        mk = GX.sniff_kind(member)
                        t = GX.document_text(member, mk)
                        if mk == "pdf" and _short(t) and len(member) > 200_000 and GX.ocr_available():
                            t = GX.ocr_pdf_text(member)
                            stat["ocr"] += 1
                            via_ocr = via_ocr or bool(t)
                        if t:
                            parts.append(t)
        except zipfile.BadZipFile:
            pass
        text = "\n".join(parts) if parts else None
        if text and via_ocr:
            text = OcrText(text)
    elif data is not None and kind == "unknown":
        stat["unknown"] += 1
    if not _short(text):
        return text
    if data is None:
        stat["dl_fail"] += 1
    elif kind == "pdf" and len(data) > 200_000:
        stat["image_pdf"] += 1
    return None


def process_board(args):
    b, days, city_obj, *rest = args
    seen_keys = rest[0] if rest else None  # 이미 처리를 끝낸 글의 키(sqlite_store.post_key). 그 글은 첨부를 다시 받지 않는다
    adapter = importlib.import_module("cms_adapters.gazette_generic")
    session = new_session()
    today = dt.date.today()
    cutoff = today - dt.timedelta(days=days - 1)
    res = {"module": b["module"], "city": b["city"], "kind": b["kind"], "url": b["url"], "status": "", "detail": "",
           "posts": 0, "attachments": 0, "notices": 0, "comp_notices": 0, "rows": 0, "ocr_rows": 0, "skipped": 0, "outcomes": [], "addr": [], "notice_list": [], "errors": []}
    board = {"name": b["kind"]}
    posts, seen, undated = [], set(), []
    last_sig = None
    try:
        for p in range(1, MAX_PAGES + 1):
            rows = adapter.fetch_list_page(session, b["url"], board, p)
            if not rows:
                break
            sig = tuple((r["title"], r["date_str"]) for r in rows[:5])
            if sig == last_sig:  # 마지막 페이지를 넘기면 같은 페이지를 되돌려주는 사이트: 끝
                break
            last_sig = sig
            older = False
            for r in rows:
                key = (r["title"], r["date_str"])
                if key in seen:
                    continue
                seen.add(key)
                if r["date"] is None:  # 날짜 없는 e-book형 목록: 최신순이라고 보고 앞쪽 몇 건만 쓴다
                    undated.append(r)
                    continue
                if r["date"] < cutoff:
                    older = True
                    continue
                posts.append(r)
            if older:
                break
            if all(r["date"] is None for r in rows):  # 날짜 없는 목록은 N일 판정이 안 되므로 첫 페이지만 쓴다
                break
            time.sleep(0.2)
    except Exception as e:
        res["status"], res["detail"] = "오류(목록)", f"{type(e).__name__}: {str(e)[:120]}"
        return res
    if not posts and undated:
        posts = undated[:UNDATED_TAKE]
        res["detail"] = f"목록에 날짜 없음 - 최신순 앞 {len(posts)}건 사용"
    if p == 1 and not posts and not seen:
        res["status"], res["detail"] = "목록 못 읽음", "날짜가 있는 목록 행을 찾지 못함(JS 화면/구조 다름)"
        return res
    res["posts"] = len(posts)
    if not posts:
        res["status"], res["detail"] = "최근 게시물 없음", f"최근 {days}일 이내 게시물 없음"
        return res
    regions = city_obj.regions if city_obj is not None else None
    # 도청·광역시 본청 공보는 여러 시군 공고가 섞여 있어 시군구를 기본값으로 채우지 않는다
    is_wide = b["city"].endswith(("도청", "광역시청", "특별자치시청")) or "통합특별시" in b["city"]
    default_sido = getattr(regions, "sido_name", "") if regions is not None and not callable(regions) else ""
    default_sigun = "" if is_wide else (getattr(regions, "sigun_name", "") if regions is not None and not callable(regions) else "")
    stat = {"hwp_only": 0, "image_pdf": 0, "dl_fail": 0, "unknown": 0, "ocr": 0, "preview": 0}
    no_detail = 0
    seen_hash = set()
    for post in posts:
        pkey = post_key(post["detail_url"] or f'{b["url"]}#{post["date_str"]}', post["title"])
        if seen_keys and pkey in seen_keys:
            res["skipped"] += 1
            continue
        err_before, rows_before, unread = len(res["errors"]), res["rows"], False
        try:
            atts = adapter.fetch_attachments(session, post["detail_url"]) if post["detail_url"] else []
        except Exception as e:
            res["errors"].append(f"상세 실패 {post['title'][:20]}: {type(e).__name__}")
            continue
        if not post["detail_url"]:
            no_detail += 1
        atts = [a for a in atts if a["ext"] in DOC_EXT and a["url"]][:MAX_ATTACH_PER_POST]
        res["attachments"] += len(atts)
        for att in atts:
            try:
                text = read_attachment_text(session, att, post["detail_url"] or b["url"], seen_hash, stat)
                if text is DUP:
                    res["attachments"] -= 1
                    continue
                if text is None:
                    unread = True  # 읽지 못한 첨부가 있으면 이 글은 끝난 게 아니다(다음에 다시 시도)
                    continue
                is_ocr = isinstance(text, OcrText)
                notices = GX.split_notices(text)
                res["notices"] += len(notices)
                for n in notices:
                    if not GX.is_compensation(n):
                        continue
                    ntitle = GX.notice_title(n)
                    if GX.is_blacklisted_title(ntitle):  # 고시공고와 같은 제외 단어(수의견적·경미한·군관리계획 등)가 제목에 있으면 수집하지 않는다
                        res["blacklisted"] = res.get("blacklisted", 0) + 1
                        continue
                    res["comp_notices"] += 1
                    rows = GX.extract_rows(n, regions, site_name=b["city"], source_file=att["name"], title=ntitle,
                                           collected_at=dt.datetime.now().isoformat(timespec="seconds"),
                                           default_sido=default_sido, default_sigun=default_sigun)
                    res["notice_list"].append({"post_title": post["title"], "post_date": post["date_str"], "post_url": post["detail_url"],
                                               "file": att["name"], "file_url": att["url"], "notice_title": ntitle, "addr_count": len(rows)})
                    for r in rows:
                        res["ocr_rows"] += int(is_ocr)
                        res["addr"].append({**r, "ocr": int(is_ocr), "post_title": post["title"], "post_date": post["date_str"], "post_url": post["detail_url"],
                                            "file_url": att["url"], "notice_title": ntitle})
                    res["rows"] += len(rows)
            except Exception as e:
                res["errors"].append(f"첨부 실패 {att['name'][:24]}: {type(e).__name__}: {str(e)[:60]}")
        if len(res["errors"]) == err_before and not unread:  # 오류 없이 끝난 글만 기록한다(오류가 난 글은 다음에 다시 읽는다)
            res["outcomes"].append({"key": pkey, "title": post["title"], "post_date": post["date_str"],
                                    "status": "ok" if res["rows"] > rows_before else "empty", "rows": res["rows"] - rows_before})
    if res["attachments"] == 0 and res["skipped"] == len(posts):
        res["status"], res["detail"] = "정상(이미 처리한 글만)", f"글 {len(posts)}건 모두 지난 실행에서 처리함"
    elif res["attachments"] == 0:
        res["status"] = "첨부 못 찾음"
        res["detail"] = "게시글은 있으나 PDF/HWP 첨부 링크를 못 찾음" + (f" (제목 링크 없음 {no_detail}건)" if no_detail else "")
    elif res["notices"] == 0 and stat["hwp_only"]:
        res["status"], res["detail"] = "HWP만 있음(미지원)", f"읽을 수 없는 .hwp(암호/배포용) {stat['hwp_only']}건"
    elif res["notices"] == 0 and stat["image_pdf"]:
        res["status"], res["detail"] = "이미지 PDF(OCR 필요)", f"OCR로도 본문을 못 읽은 스캔형 PDF {stat['image_pdf']}건"
    elif res["notices"] == 0 and stat["dl_fail"]:
        res["status"], res["detail"] = "다운로드 실패", f"첨부 {stat['dl_fail']}건 다운로드 실패(미리보기도 안 열림)"
    elif res["notices"] == 0:
        res["status"], res["detail"] = "본문 추출 실패", "첨부는 받았으나 텍스트를 못 읽음"
    elif res["comp_notices"] == 0:
        res["status"], res["detail"] = "정상(보상 공고 없음)", f"공고 {res['notices']}건 중 보상 관련 없음"
    elif res["rows"] == 0:
        res["status"], res["detail"] = "정상(주소 0건)", f"보상 공고 {res['comp_notices']}건, 주소 표기를 못 찾음"
    else:
        res["status"] = "정상"
        res["detail"] = f"보상 공고 {res['comp_notices']}건 / 주소 {res['rows']}건"
    if stat["ocr"] or stat["preview"]:
        res["detail"] += f" | OCR {stat['ocr']}건, 미리보기 {stat['preview']}건, OCR 출처 행 {res['ocr_rows']}건"
    if res["skipped"]:
        res["detail"] += f" | 이미 처리한 글 {res['skipped']}건 건너뜀"
    if res.get("blacklisted"):
        res["detail"] += f" | 제외 단어 공고 {res['blacklisted']}건 건너뜀"
    if res["errors"]:
        res["detail"] += f" | 오류 {len(res['errors'])}건"
    return res


def main():
    ap = argparse.ArgumentParser(description="지자체 공보 보상 주소 추출")
    ap.add_argument("--excel", default=DEFAULT_EXCEL)
    ap.add_argument("--days", type=int, default=15)
    ap.add_argument("--city", action="append")
    ap.add_argument("--module", action="append", help="권역(파일) 단위로 선택: 예) 경기, 강원, 충북")
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--output", default=os.path.join(HERE, "output", "gazette"))
    a = ap.parse_args()

    boards = load_boards(a.excel)
    if a.city:
        boards = [b for b in boards if b["city"] in a.city]
    if a.module:
        boards = [b for b in boards if b["module"] in a.module]
    idx = regions_index()
    jobs = [(b, a.days, idx.get((b["module"], b["city"]))) for b in boards]
    print(f"대상 게시판 {len(jobs)}곳, 최근 {a.days}일")
    t0 = time.time()
    with ThreadPoolExecutor(a.workers) as ex:
        results = list(ex.map(process_board, jobs))
    out = Path(a.output)
    out.mkdir(parents=True, exist_ok=True)
    done = {(r["module"], r["city"]) for r in results}
    addr_fields = ["module", "city", "kind", "post_date", "post_title", "notice_title", "sojaeji", "jibun", "jimok",
                   "area_incl", "project_name", "remark", "post_url", "file_url", "source_file", "collected_at"]
    notice_fields = ["module", "city", "kind", "post_date", "post_title", "notice_title", "addr_count", "file", "post_url", "file_url"]

    def merge_csv(path, fields, new_rows):
        """이미 있는 결과 중 이번에 다시 처리한 지자체 행은 버리고 새 행을 붙인다(권역별로 나눠 실행해도 누적)."""
        old_rows = []
        if path.exists():
            with open(path, encoding="utf-8-sig", newline="") as f:
                old_rows = [r for r in csv.DictReader(f) if (r["module"], r["city"]) not in done]
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader()
            for r in old_rows + new_rows:
                w.writerow(r)

    merge_csv(out / "gazette_addresses.csv", addr_fields,
              [{"module": r["module"], "city": r["city"], "kind": r["kind"], **x} for r in results for x in r["addr"]])
    merge_csv(out / "gazette_notices.csv", notice_fields,
              [{"module": r["module"], "city": r["city"], "kind": r["kind"], **x} for r in results for x in r["notice_list"]])
    slim = [{k: v for k, v in r.items() if k not in ("addr", "notice_list")} for r in results]
    st_path = out / "gazette_status.json"
    old_st = json.load(open(st_path, encoding="utf-8")) if st_path.exists() else []
    old_st = [x for x in old_st if (x["module"], x["city"]) not in done]
    json.dump(old_st + slim, open(st_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    from collections import Counter
    print("결과:", dict(Counter(r["status"] for r in results)), f"| 주소 {sum(r['rows'] for r in results)}건 | {time.time()-t0:.0f}초")


if __name__ == "__main__":
    main()
